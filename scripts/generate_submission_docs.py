"""Generate professional submission artifacts for Snapdragon AI Lab Challenge.

Produces:
  - submission/Aegis_Project_Description.docx
  - submission/Aegis_Project_Description.pdf
  - submission/Aegis_Pitch_Deck.pptx
  - submission/Aegis_Pitch_Deck.pdf
"""

from __future__ import annotations

import os
import sys
from pathlib import Path

# Paths
PROJECT_ROOT = Path(__file__).resolve().parent.parent
SUBMISSION_DIR = PROJECT_ROOT / "submission"
SUBMISSION_DIR.mkdir(exist_ok=True)
ASSETS_DIR = PROJECT_ROOT / "docs" / "assets"
PREVIEW_IMG = ASSETS_DIR / "aegis-core-dashboard.png"

# ==============================================================================
# 1. GENERATE PROJECT DESCRIPTION (.docx & .pdf)
# ==============================================================================

def generate_project_description_docx(output_path: Path) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn

    doc = Document()

    # Configure Margins (0.75 in)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    def set_cell_background(cell, fill_hex):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{m}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    NAVY = RGBColor(11, 37, 69)     # #0B2545
    CYAN = RGBColor(0, 140, 186)    # #008CBA
    DARK = RGBColor(33, 37, 41)     # #212529
    GREY = RGBColor(108, 117, 125)  # #6C757D

    # Title Block
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("Aegis: Privacy-Preserving On-Device Combat Motion Intelligence")
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = NAVY

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("Snapdragon AI Lab Build & Present Challenge Submission  |  Detailed Technical Specification")
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = CYAN

    # Metadata Table
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    meta_data = [
        [("Participant", "Aniket Kapgate"), ("Target Platform", "Snapdragon X Elite / Plus HP PCs (Windows 11 on Arm)")],
        [("Repository", "https://github.com/aniketkapgate7-crypto/Aegis-Fight-Pattern-Analyzer"), ("Evaluation Status", "Verified MediaPipe CPU Baseline | QNN Optimization Planned")]
    ]
    for r_idx, row in enumerate(meta_data):
        for c_idx, (label, val) in enumerate(row):
            cell = meta_table.cell(r_idx, c_idx)
            set_cell_background(cell, "F4F7FA")
            set_cell_margins(cell, 80, 80, 120, 120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r_lbl = p.add_run(f"{label}: ")
            r_lbl.font.bold = True
            r_lbl.font.size = Pt(9.5)
            r_lbl.font.color.rgb = NAVY
            r_val = p.add_run(val)
            r_val.font.size = Pt(9.5)
            r_val.font.color.rgb = DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Styling Helpers
    def add_heading(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.font.name = "Segoe UI"
        r.font.size = Pt(13 if level == 1 else 11)
        r.font.bold = True
        r.font.color.rgb = NAVY
        return p

    def add_body(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.font.name = "Segoe UI"
        r.font.size = Pt(10)
        r.font.color.rgb = DARK
        return p

    def add_bullet(text, prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        if prefix:
            r_pre = p.add_run(prefix)
            r_pre.font.name = "Segoe UI"
            r_pre.font.size = Pt(10)
            r_pre.font.bold = True
            r_pre.font.color.rgb = NAVY
        r = p.add_run(text)
        r.font.name = "Segoe UI"
        r.font.size = Pt(10)
        r.font.color.rgb = DARK
        return p

    # 1. EXECUTIVE SUMMARY
    add_heading("1. Executive Summary")
    add_body(
        "Aegis is an edge-native, privacy-preserving motion intelligence system engineered for real-time human movement "
        "and threat pattern analysis on Snapdragon-powered HP PCs. The system transforms live optical camera feeds into "
        "normalized 2D skeletal landmarks, evaluates multi-frame kinematic trajectories over a temporal sliding window, "
        "and identifies strikes, defensive guards, rapid approach, and fall patterns. Crucially, Aegis processes video "
        "entirely in local memory without transmitting imagery to cloud servers, stores no biometric templates, and does "
        "not perform facial recognition. Aegis does not perform face recognition, face identification, facial embeddings or "
        "biometric identity tracking. One coarse nose keypoint is used only for pose geometry. By uniting zero-cloud on-device "
        "inference with explainable kinematic evidence, Aegis provides transparent decision-support telemetry for sports coaching, "
        "supervised athletic safety, and edge AI research."
    )

    # 2. PROBLEM STATEMENT
    add_heading("2. Problem Statement")
    add_body(
        "Contemporary physical safety monitoring and athletic combat analysis rely almost universally on two flawed paradigms:"
    )
    add_bullet(
        " Traditional video surveillance and coaching reviews occur post-incident, failing to alert supervisors during the unfolding event.",
        "Reactive Post-Event Review:"
    )
    add_bullet(
        " Offloading video streams to cloud analytics introduces substantial transmission latency (200-800+ ms), consumes critical network bandwidth, and fails when Internet connectivity drops.",
        "Cloud Latency & Network Fragility:"
    )
    add_bullet(
        " Transmitting unencrypted visual feeds of athletes or public spaces to remote datacenters creates catastrophic risks of biometric data theft, facial surveillance abuses, and non-consensual tracking.",
        "Severe Biometric & Privacy Exposure:"
    )
    add_bullet(
        " End-to-end deep video classifiers output opaque probability vectors without explaining which joint, limb velocity, or physical movement triggered the classification.",
        "Opaque Black-Box Classifications:"
    )

    # 3. PROPOSED SOLUTION
    add_heading("3. Proposed Solution")
    add_body(
        "Aegis addresses these core vulnerabilities through a four-pillar edge architecture:"
    )
    add_bullet(
        " Ingests camera streams into volatile RAM, calculates joint coordinates, and immediately overwrites frame buffers. No video is uploaded to cloud endpoints.",
        "100% On-Device Processing:"
    )
    add_bullet(
        " Completely abstracts human subjects into 9 skeletal joint coordinates (nose, shoulders, wrists, hips, ankles), stripping clothing, ethnicity, and identity. Aegis does not perform face recognition, face identification, facial embeddings or biometric identity tracking. One coarse nose keypoint is used only for pose geometry.",
        "Biometric Anonymity:"
    )
    add_bullet(
        " Evaluates physical kinematics (limb reach multiples, angular velocities, torso ratios) over a 12-frame sliding window, producing human-verifiable evidence strings.",
        "Explainable Temporal Logic:"
    )
    add_bullet(
        " Renders a high-contrast 960x540 heads-up display with live FPS, inference latency, threat gauges, and optional deduplicated local incident logging.",
        "AEGIS CORE Telemetry:"
    )

    # Insert Image
    if PREVIEW_IMG.exists():
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(2)
        doc.add_picture(str(PREVIEW_IMG), width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(10)
        r_cap = p_cap.add_run("Figure 1: AEGIS CORE live heads-up display operating at 960x540 with synthetic martial arts pose telemetry.")
        r_cap.font.name = "Segoe UI"
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = GREY

    # 4. TARGET USERS & USE CASES
    add_heading("4. Target Users and Use Cases")
    add_bullet(
        " Quantitative velocity, extension reach, and guard recovery metrics for boxing, karate, and mixed martial arts training without bulky wearable sensors.",
        "Athletic & Combat Sports Coaching:"
    )
    add_bullet(
        " Real-time awareness for facility supervisors monitoring wrestling mats, gymnastics arenas, and athletic centers to detect sudden falls or physical distress.",
        "Supervised Facility Safety:"
    )
    add_bullet(
        " Demonstrating how Qualcomm Snapdragon Hexagon NPUs can offload computer vision pipelines on battery-constrained laptops without sacrificing privacy.",
        "Edge AI & NPU Research:"
    )

    # 5. SYSTEM ARCHITECTURE & TECHNICAL IMPLEMENTATION
    add_heading("5. System Architecture & Technical Implementation")
    add_body(
        "Aegis is constructed from decoupled, modular subsystems ensuring clean testability, zero global state, and straightforward hardware acceleration migration:"
    )
    add_bullet(
        " Ingests live webcam feeds or pre-recorded clips via OpenCV, downsampling to 960x540 for optimal balance between landmark resolution and display latency.",
        "Frame Ingestion & Preprocessing:"
    )
    add_bullet(
        " Extracts 9 primary joint landmarks with individual visibility confidence gating (>0.40) to prevent erroneous tracking under limb occlusion.",
        "Pose Estimation Engine (MediaPipe CPU):"
    )
    add_bullet(
        " Maintains a deque of recent poses (12 frames) to calculate velocity, body scale expansion, and positional differentials.",
        "Temporal Feature Window:"
    )
    add_bullet(
        " Evaluates joint geometry against validated biomechanical thresholds with strict priority hierarchy (Fall > Punch > Kick > Approach > Guard > Neutral).",
        "Kinematic Pattern Engine:"
    )
    add_bullet(
        " Aggregates confidence and threat weights into a unified [0.0, 1.0] scale, mapped to LOW, ELEVATED, HIGH, and CRITICAL danger states.",
        "Threat Fusion Engine:"
    )
    add_bullet(
        " Renders custom anti-aliased HUD panels, threat gauges, and timelines. Manages local JSONL incident logging with state-driven threat re-arming.",
        "AEGIS CORE UI & Incident Tracker:"
    )

    # 6. PATTERN-DETECTION METHODOLOGY
    add_heading("6. Pattern-Detection Methodology")
    add_body("The pattern recognition engine evaluates six distinct movement classes deterministically:")

    pattern_table = doc.add_table(rows=7, cols=4)
    pattern_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Pattern", "Kinematic Criteria", "Threat", "Sample Rationale"]
    col_widths = [Inches(1.4), Inches(2.2), Inches(0.8), Inches(2.1)]
    for i, h in enumerate(headers):
        cell = pattern_table.cell(0, i)
        set_cell_background(cell, "0B2545")
        p = cell.paragraphs[0]
        r = p.add_run(h)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9)

    p_rows = [
        ("Neutral", "Upright torso, limbs in equilibrium", "0.05", "No risk-pattern threshold crossed"),
        ("Guard", "Wrists within 1.30x shoulder width of nose", "0.30", "Both hands raised near face"),
        ("Punch-like Extension", "Arm reach > 1.45x shoulder width, speed > 0.90/s", "0.88", "Right arm rapidly extended (reach 1.65x)"),
        ("Kick-like Extension", "Ankle reach > 1.35x, speed > 0.65/s, elevated", "0.90", "Right leg extended, ankle elevated"),
        ("Rapid Approach", "Torso scale growth > 6% over 3 frames, rate > 0.80/s", "0.72", "Body scale increased rapidly (+12.4%)"),
        ("Fall-like Posture", "Torso horizontal/vertical ratio > 1.25, compact", "0.90", "Torso horizontal ratio 1.84, vertical span low"),
    ]

    for r_idx, row_data in enumerate(p_rows):
        for c_idx, text in enumerate(row_data):
            cell = pattern_table.cell(r_idx + 1, c_idx)
            set_cell_background(cell, "F9FAFC" if r_idx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, 60, 60, 80, 80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(8.5)
            r.font.color.rgb = DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # 7. EXPLAINABILITY & HUMAN-IN-THE-LOOP TELEMETRY
    add_heading("7. Explainability & Human-in-the-Loop Telemetry")
    add_body(
        "A critical limitation of deep learning in safety-critical domains is the 'black box' problem. Aegis ensures complete "
        "interpretability: every incident record and HUD alert includes human-readable telemetry explaining exactly why the alert "
        "triggered (e.g., 'right arm reach 1.62x shoulder width, extension speed 1.15/s'). Aegis is strictly designed as an operator "
        "decision-support tool; autonomous physical locking, punitive actions, or automated enforcement are prohibited."
    )

    # 8. PRIVACY & RESPONSIBLE-USE SAFEGUARDS
    add_heading("8. Privacy & Responsible-Use Safeguards")
    add_bullet(" Video is processed in volatile memory. No raw video frames are saved to disk or transmitted across network interfaces.", "Zero Cloud Upload:")
    add_bullet(" Aegis does not perform face recognition, face identification, facial embeddings or biometric identity tracking. One coarse nose keypoint is used only for pose geometry.", "Biometric Anonymity:")
    add_bullet(" Logging occurs only when an operator explicitly toggles recording (R or Space). Log files reside locally in artifacts/.", "Opt-In Local Logging:")
    add_bullet(" Aegis detects physical motions, not criminal intent or moral guilt. Human verification is strictly required.", "No Intent Inference:")

    # 9. CURRENT VERIFIED EVIDENCE
    add_heading("9. Current Verified Evidence & Technical Truthfulness")
    add_body(
        "In strict accordance with competition rules, all claims in Aegis reflect verified, reproducible execution:"
    )
    add_bullet(" Fully operational video capture, landmark ingestion, and 960x540 display rendering.", "Functional Live Pipeline:")
    add_bullet(" 17 deterministic unit tests passing in pytest (covering neutral, strikes, guard, falls, approach, jitter rejection, boundaries, deduplication, and provider reporting).", "Automated Test Suite:")
    add_bullet(" MediaPipe CPU is the only currently verified active pose provider. Supplying AEGIS_MODEL_PATH does not connect model outputs to pose estimation; runtime diagnostics only verify provider availability. Genuine QNN pose execution requires an ONNX/QNN pose-estimator adapter that preprocesses frames, calls session.run(), converts outputs to normalized PoseFrame landmarks, and replaces MediaPipePoseEstimator.", "Truthful Provider Reporting:")
    add_bullet(" Re-arming mechanism validated; prevents duplicate incident records during continuous high-threat motions.", "Deduplication Verification:")

    # 10. SNAPDRAGON OPTIMIZATION PLAN
    add_heading("10. Snapdragon Optimization Plan & Qualcomm AI Hub (Planned)")
    add_body(
        "Aegis features a modular runtime interface designed for transition to Qualcomm Hexagon NPUs on Snapdragon-powered HP PCs. "
        "Supplying AEGIS_MODEL_PATH does not connect model outputs to pose estimation; runtime diagnostics only verify provider availability. "
        "MediaPipe CPU is the only currently verified active pose provider. The planned optimization pathway includes:"
    )
    add_bullet(" Select a lightweight pose estimation network (YOLOv8n-pose or RTMPose) from the Qualcomm AI Hub catalog.", "1. Model Selection (Planned):")
    add_bullet(" Quantize (INT8) and compile the model specifically targeting the Snapdragon X Elite / Plus Hexagon NPU using the Qualcomm AI Hub compilation API.", "2. Model Compilation (Planned):")
    add_bullet(" Implement an ONNX/QNN pose-estimator adapter that preprocesses frames, calls session.run(), converts outputs to normalized PoseFrame landmarks, and replaces MediaPipePoseEstimator via ONNX Runtime QNNExecutionProvider (QnnHtp.dll).", "3. NPU Execution (Planned):")
    add_bullet(" Execute scripts/benchmark.py on an identical 1080p clip to measure end-to-end FPS, p95 latency, and power savings.", "4. Verification Protocol (Planned):")

    # 11. ACCESSIBILITY & INNOVATION
    add_heading("11. Accessibility & Key Innovations")
    add_bullet(" High-contrast color palette, readable text shadows, combined geometric icons, numeric readouts, and clear text status labels.", "Accessible HUD:")
    add_bullet(" Full keyboard control (R/Space for recording, Q/Esc for exit) and deterministic synthetic mode (--demo / --export-preview).", "Inclusive Controls:")
    add_bullet(" Temporal biomechanical analysis rather than single-frame classification; transparent evidence generation; truthful on-device runtime abstraction.", "Key Innovations:")

    # 12. LIMITATIONS & ROADMAP
    add_heading("12. Limitations & Future Roadmap")
    add_bullet(" Extreme camera angles (steep overhead) and severe visual occlusion degrade 2D coordinate accuracy, triggering safe neutral fallback.", "Known Limitations:")
    add_bullet(" QNN Hexagon NPU pose integration, multi-person tracking support, and battery-efficiency benchmarking on Snapdragon X Elite HP laptops.", "Future Roadmap:")

    doc.save(str(output_path))
    print(f"Project Description DOCX generated at: {output_path.resolve()}")


def generate_project_description_pdf(docx_path: Path, output_path: Path) -> None:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.pdfgen import canvas

    class NumberedCanvas(canvas.Canvas):
        def __init__(self, *args, **kwargs):
            canvas.Canvas.__init__(self, *args, **kwargs)
            self._saved_page_states = []

        def showPage(self):
            self._saved_page_states.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            num_pages = len(self._saved_page_states)
            for state in self._saved_page_states:
                self.__dict__.update(state)
                self.draw_page_number(num_pages)
                canvas.Canvas.showPage(self)
            canvas.Canvas.save(self)

        def draw_page_number(self, page_count):
            self.saveState()
            self.setFont("Helvetica", 8.5)
            self.setFillColor(colors.HexColor("#6C757D"))
            # Header
            self.drawString(54, 750, "Aegis: Privacy-Preserving On-Device Combat Motion Intelligence  |  Challenge Specification")
            self.setStrokeColor(colors.HexColor("#D1D5DB"))
            self.setLineWidth(0.5)
            self.line(54, 744, 558, 744)
            # Footer
            self.line(54, 45, 558, 45)
            page_text = f"Page {self._pageNumber} of {page_count}"
            self.drawRightString(558, 32, page_text)
            self.drawString(54, 32, "Snapdragon AI Lab Challenge  |  Participant: Aniket Kapgate  |  CONFIDENTIAL")
            self.restoreState()

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0B2545"),
        spaceAfter=4,
    )
    sub_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#008CBA"),
        spaceAfter=10,
    )
    h1_style = ParagraphStyle(
        'Heading1_Custom',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=12.5,
        leading=16,
        textColor=colors.HexColor("#0B2545"),
        spaceBefore=10,
        spaceAfter=4,
        keepWithNext=True,
    )
    body_style = ParagraphStyle(
        'Body_Custom',
        parent=styles['BodyText'],
        fontName='Helvetica',
        fontSize=9.2,
        leading=13.2,
        textColor=colors.HexColor("#212529"),
        spaceAfter=5,
    )
    bullet_style = ParagraphStyle(
        'Bullet_Custom',
        parent=styles['BodyText'],
        fontName='Helvetica',
        fontSize=9.2,
        leading=13.2,
        textColor=colors.HexColor("#212529"),
        leftIndent=15,
        firstLineIndent=-10,
        spaceAfter=3,
    )

    story = []

    # Title & Subtitle
    story.append(Paragraph("Aegis: Privacy-Preserving On-Device Combat Motion Intelligence", title_style))
    story.append(Paragraph("Snapdragon AI Lab Build & Present Challenge Submission  |  Technical Description", sub_style))

    # Meta Table
    meta_data = [
        [Paragraph("<b>Participant:</b> Aniket Kapgate", body_style), Paragraph("<b>Target Platform:</b> Snapdragon X Elite / Plus HP PCs", body_style)],
        [Paragraph("<b>Repository:</b> github.com/aniketkapgate7-crypto/Aegis-Fight-Pattern-Analyzer", body_style), Paragraph("<b>Status:</b> Verified MediaPipe CPU Baseline | QNN Optimization Planned", body_style)]
    ]
    t_meta = Table(meta_data, colWidths=[250, 254])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F4F7FA")),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(t_meta)
    story.append(Spacer(1, 8))

    # 1. Executive Summary
    story.append(Paragraph("1. Executive Summary", h1_style))
    story.append(Paragraph(
        "<b>Aegis</b> is an on-device, privacy-preserving motion intelligence application designed to deliver real-time, explainable combat and athletic movement analysis on Snapdragon-powered HP PCs. By combining normalized skeletal pose landmarks with a sliding-window temporal kinematic engine, Aegis detects rapid strikes, defensive guards, sudden approach, and fall patterns directly on the local workstation. It operates completely offline, eliminates cloud video streaming vulnerabilities, and empowers human operators with actionable, explainable telemetry. Aegis does not perform face recognition, face identification, facial embeddings or biometric identity tracking. One coarse nose keypoint is used only for pose geometry.",
        body_style
    ))

    # 2. Problem Statement
    story.append(Paragraph("2. Problem Statement", h1_style))
    story.append(Paragraph("• <b>Reactive Post-Incident Review:</b> Incidents and athletic form breakdowns are typically reviewed hours after occurrence from passive recordings, failing to prevent injury.", bullet_style))
    story.append(Paragraph("• <b>Cloud Latency & Network Fragility:</b> Streaming raw video to cloud servers introduces latency (200-800ms), consumes network bandwidth, and fails entirely during connectivity interruptions.", bullet_style))
    story.append(Paragraph("• <b>Biometric & Privacy Vulnerabilities:</b> Transmitting video over public networks exposes individuals to facial recognition, identity harvesting, and severe privacy violations.", bullet_style))
    story.append(Paragraph("• <b>Opaque Classification:</b> Conventional vision models output unexplainable classification scores without identifying which limb or motion triggered the alert.", bullet_style))

    # 3. Proposed Solution
    story.append(Paragraph("3. Proposed Solution", h1_style))
    story.append(Paragraph(
        "Aegis resolves these issues through an edge-native, human-in-the-loop movement intelligence pipeline: "
        "<b>100% On-Device Processing</b> processes frames in local RAM without cloud transmission. "
        "<b>Skeletal Landmark Abstraction</b> converts pixels into 9 normalized joint coordinates. Aegis does not perform face recognition, face identification, facial embeddings or biometric identity tracking. One coarse nose keypoint is used only for pose geometry. "
        "<b>Explainable Kinematics</b> evaluates multi-frame limb reach and velocity. "
        "<b>AEGIS CORE HUD</b> displays live FPS, latency, threat meters, and optional deduplicated incident logs.",
        body_style
    ))

    # Preview Image
    if PREVIEW_IMG.exists():
        story.append(Spacer(1, 4))
        story.append(Image(str(PREVIEW_IMG), width=5.8*inch, height=3.26*inch))
        cap_style = ParagraphStyle('Cap', parent=body_style, fontSize=8, leading=10, textColor=colors.HexColor("#6C757D"), alignment=1)
        story.append(Paragraph("Figure 1: AEGIS CORE live heads-up display operating at 960x540 with synthetic martial arts pose telemetry.", cap_style))
        story.append(Spacer(1, 6))

    # 4. Target Users & Use Cases
    story.append(Paragraph("4. Target Users & Use Cases", h1_style))
    story.append(Paragraph("• <b>Combat Sports & Athletic Coaching:</b> Quantitative feedback on punch extension speed, guard discipline, and kick reach in boxing and martial arts training.", bullet_style))
    story.append(Paragraph("• <b>Supervised Facility Safety:</b> Assisting athletic supervisors in monitoring training mats to detect sudden falls, physical trauma, or aggressive collisions with human-in-the-loop oversight.", bullet_style))
    story.append(Paragraph("• <b>Edge AI & NPU Research:</b> Demonstrating how Snapdragon Hexagon NPUs can offload computer vision tasks from host CPU cores efficiently.", bullet_style))

    # 5. System Architecture
    story.append(Paragraph("5. System Architecture & Technical Implementation", h1_style))
    story.append(Paragraph(
        "Aegis employs a strictly decoupled, modular architecture: "
        "<b>Frame Acquisition:</b> OpenCV captures video stream, resizing to 960x540. "
        "<b>Pose Estimator:</b> MediaPipe Pose generates 9 keypoints with visibility gating (>0.40). "
        "<b>Temporal Window:</b> A 12-frame deque tracks joint coordinates over time. "
        "<b>Pattern Engine:</b> Evaluates kinematic rules (Fall > Punch > Kick > Approach > Guard > Neutral). "
        "<b>Threat Fusion:</b> Combines signals into a 0.0-1.0 threat index. "
        "<b>AEGIS CORE UI:</b> Renders HUD and manages deduplicated JSONL incident logging.",
        body_style
    ))

    # 6. Pattern Methodology Table
    story.append(Paragraph("6. Pattern-Detection Methodology", h1_style))
    p_data = [
        [Paragraph("<b>Pattern</b>", body_style), Paragraph("<b>Kinematic Criteria</b>", body_style), Paragraph("<b>Threat</b>", body_style), Paragraph("<b>Evidence String</b>", body_style)],
        [Paragraph("Neutral", body_style), Paragraph("Equilibrium, no thresholds crossed", body_style), Paragraph("0.05", body_style), Paragraph("no risk-pattern threshold crossed", body_style)],
        [Paragraph("Guard", body_style), Paragraph("Wrists within 1.30x shoulder width of nose", body_style), Paragraph("0.30", body_style), Paragraph("both hands raised near face", body_style)],
        [Paragraph("Punch", body_style), Paragraph("Arm reach >1.45x, speed >0.90/s", body_style), Paragraph("0.88", body_style), Paragraph("right arm rapidly extended (reach 1.65x)", body_style)],
        [Paragraph("Kick", body_style), Paragraph("Leg reach >1.35x, speed >0.65/s, elevated", body_style), Paragraph("0.90", body_style), Paragraph("right leg extended, ankle elevated", body_style)],
        [Paragraph("Rapid Approach", body_style), Paragraph("Torso scale growth >6% across 3 frames", body_style), Paragraph("0.72", body_style), Paragraph("body scale increased rapidly (+12.4%)", body_style)],
        [Paragraph("Fall", body_style), Paragraph("Torso horizontal ratio >1.25, compact", body_style), Paragraph("0.90", body_style), Paragraph("torso horizontal ratio 1.84", body_style)],
    ]
    t_pattern = Table(p_data, colWidths=[85, 175, 45, 199])
    t_pattern.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0B2545")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t_pattern)
    story.append(Spacer(1, 6))

    # 7. Privacy & Responsible Use
    story.append(Paragraph("7. Privacy, Safeguards & Responsible-Use", h1_style))
    story.append(Paragraph("• <b>Zero Cloud Streaming:</b> Video is processed in volatile memory and immediately overwritten.", bullet_style))
    story.append(Paragraph("• <b>No Facial Biometrics:</b> Aegis does not perform face recognition, face identification, facial embeddings or biometric identity tracking. One coarse nose keypoint is used only for pose geometry.", bullet_style))
    story.append(Paragraph("• <b>Opt-In Local Logging:</b> Logging occurs only when enabled by the operator (R or Space key).", bullet_style))
    story.append(Paragraph("• <b>Human Interpretation Required:</b> Aegis outputs motion metrics, not intent or guilt. Autonomous enforcement is strictly prohibited.", bullet_style))

    # 8. Evidence & Snapdragon Optimization
    story.append(Paragraph("8. Verified Evidence & Snapdragon Optimization Plan", h1_style))
    story.append(Paragraph(
        "<b>Current Verified Baseline:</b> Functional webcam ingestion, 17/17 deterministic unit tests passing in pytest, "
        "and truthful provider reporting (MediaPipe CPU is the only currently verified active pose provider). Supplying AEGIS_MODEL_PATH does not connect model outputs to pose estimation; runtime diagnostics only verify provider availability. "
        "<b>Planned Snapdragon Migration Pathway:</b> 1) Select compatible pose model (YOLOv8n-pose or RTMPose) from Qualcomm AI Hub. "
        "2) Compile and quantize (INT8) targeting Hexagon NPU on Snapdragon X Elite/Plus HP laptops. "
        "3) Genuine QNN pose execution requires an ONNX/QNN pose-estimator adapter that preprocesses frames, calls session.run(), converts outputs to normalized PoseFrame landmarks, and replaces MediaPipePoseEstimator via QNNExecutionProvider. "
        "4) Run scripts/benchmark.py on an identical clip to capture comparative FPS and p95 latency gains.",
        body_style
    ))

    # 9. GitHub Link & Status
    story.append(Paragraph("9. Repository & Challenge Status", h1_style))
    story.append(Paragraph(
        "<b>GitHub Repository:</b> https://github.com/aniketkapgate7-crypto/Aegis-Fight-Pattern-Analyzer<br/>"
        "<b>Branch:</b> feat/snapdragon-submission-readiness  |  <b>Challenge Track:</b> Build & Present",
        body_style
    ))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"Project Description PDF generated at: {output_path.resolve()}")


# ==============================================================================
# 2. GENERATE SEVEN-SLIDE PITCH DECK (.pptx & .pdf)
# ==============================================================================

def generate_pitch_deck_pptx(output_path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette (Polished Dark Navy / Black Theme with Cyan Accents)
    BG_COLOR = RGBColor(7, 15, 20)          # #070F14 (Deep Tech Dark)
    PANEL_BG = RGBColor(14, 30, 40)         # #0E1E28 (Translucent Card)
    PANEL_BORDER = RGBColor(28, 58, 77)     # #1C3A4D (Card Border)
    CYAN = RGBColor(0, 229, 255)            # #00E5FF (High-Tech Cyan)
    CYAN_DIM = RGBColor(0, 140, 186)        # #008CBA
    WHITE = RGBColor(245, 248, 250)         # #F5F8FA
    GREY = RGBColor(165, 180, 188)          # #A5B4BC
    GREEN = RGBColor(70, 245, 85)           # #46F555
    ORANGE = RGBColor(255, 185, 0)          # #FFB900
    RED = RGBColor(245, 55, 40)             # #F53728

    def add_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()
        return bg

    def add_header(slide, category: str, title: str):
        # Category indicator
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.4))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = f"/// {category.upper()}"
        p_cat.font.name = "Segoe UI"
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = CYAN

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.name = "Segoe UI"
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

        # Horizontal accent rule
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.65), Inches(11.733), Inches(0.02))
        rule.fill.solid()
        rule.fill.fore_color.rgb = PANEL_BORDER
        rule.line.fill.background()

    def add_card(slide, x, y, w, h, title="", subtitle=""):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        card.fill.solid()
        card.fill.fore_color.rgb = PANEL_BG
        card.line.color.rgb = PANEL_BORDER
        card.line.width = Pt(1)

        if title:
            tbox = slide.shapes.add_textbox(x + Inches(0.2), y + Inches(0.15), w - Inches(0.4), Inches(0.4))
            tf = tbox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.name = "Segoe UI"
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = CYAN

        if subtitle:
            sbox = slide.shapes.add_textbox(x + Inches(0.2), y + Inches(0.55), w - Inches(0.4), h - Inches(0.7))
            tf = sbox.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = subtitle
            p.font.name = "Segoe UI"
            p.font.size = Pt(10)
            p.font.color.rgb = GREY

        return card

    # --------------------------------------------------------------------------
    # SLIDE 1 — TITLE SLIDE
    # --------------------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    add_background(s1)

    # Accent decorative lines
    dec1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.8), Inches(0.08), Inches(3.6))
    dec1.fill.solid()
    dec1.fill.fore_color.rgb = CYAN
    dec1.line.fill.background()

    t_box = s1.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.5), Inches(3.6))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    p1.text = "AEGIS"
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(44)
    p1.font.bold = True
    p1.font.color.rgb = CYAN

    p2 = tf1.add_paragraph()
    p2.text = "Privacy-Preserving On-Device Combat Motion Intelligence"
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.space_after = Pt(14)

    p3 = tf1.add_paragraph()
    p3.text = "Explainable real-time motion awareness designed for Snapdragon-powered HP PCs"
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(15)
    p3.font.color.rgb = GREY
    p3.space_after = Pt(24)

    p4 = tf1.add_paragraph()
    p4.text = "Participant: Aniket Kapgate  |  Snapdragon AI Lab Build & Present Challenge"
    p4.font.name = "Segoe UI"
    p4.font.size = Pt(12)
    p4.font.bold = True
    p4.font.color.rgb = CYAN_DIM

    # --------------------------------------------------------------------------
    # SLIDE 2 — PROBLEM
    # --------------------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_background(s2)
    add_header(s2, "Challenge & Market Need", "The Problem: Latency, Fragility, and Surveillance Exposure")

    problems = [
        ("01 / REACTIVE REVIEW", "Incident Review Happens After the Fact", "In athletic coaching and facility safety, video is examined hours or days later. Supervisors lack real-time warning to intervene before injuries or trauma occur."),
        ("02 / CLOUD DEPENDENCE", "High Latency & Single Point of Failure", "Cloud video analytics streams high-bandwidth video over public internet, introducing 200-800ms delays and completely collapsing during network outages."),
        ("03 / TEMPORAL BLINDNESS", "Single-Frame Models Miss Motion Context", "Static single-frame object detectors cannot distinguish a high-velocity punch from an open high-five or a stumble from a dynamic athletic slide."),
        ("04 / PRIVACY VULNERABILITY", "Biometric Harvesting & Facial Exposure", "Streaming surveillance footage to external cloud servers creates catastrophic risks of non-consensual biometric data scraping and surveillance abuses.")
    ]

    card_w = Inches(2.75)
    card_gap = Inches(0.24)
    start_x = Inches(0.8)
    for idx, (tag, ptitle, pdesc) in enumerate(problems):
        cx = start_x + idx * (card_w + card_gap)
        add_card(s2, cx, Inches(2.1), card_w, Inches(4.6), tag, "")
        
        # Add detailed text inside card
        tb = s2.shapes.add_textbox(cx + Inches(0.2), Inches(2.7), card_w - Inches(0.4), Inches(3.8))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = ptitle
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(13)
        pt.font.bold = True
        pt.font.color.rgb = WHITE
        pt.space_after = Pt(12)
        
        pd = tf.add_paragraph()
        pd.text = pdesc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(10)
        pd.font.color.rgb = GREY

    # --------------------------------------------------------------------------
    # SLIDE 3 — SOLUTION
    # --------------------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_background(s3)
    add_header(s3, "Product Vision", "The Solution: Edge-Native Biomechanical Intelligence")

    sol_cards = [
        ("LOCAL PROCESSING", "100% On-Device Execution", "Processes video frames entirely in local RAM on Snapdragon PCs. Video buffers are analyzed and immediately overwritten—never uploaded."),
        ("ANONYMIZED SKELETONS", "Human Pose Landmarks", "Extracts 9 normalized 2D skeletal landmarks (wrists, shoulders, hips, ankles). Aegis does not perform face recognition, face identification, facial embeddings or biometric identity tracking. One coarse nose keypoint is used only for pose geometry."),
        ("TEMPORAL WINDOWS", "Multi-Frame Motion Context", "Maintains a 12-frame sliding window to calculate true physical velocities, body scale growth rates, and torso aspect ratios."),
        ("ACTION VOCABULARY", "6 Biomechanical Patterns", "Deterministically detects Neutral, Guard, Punch-like extension, Kick-like extension, Rapid Approach, and Fall-like recumbent postures."),
        ("TRANSPARENT METRICS", "Explainable Telemetry", "Outputs human-verifiable evidence strings (e.g., 'right arm reach 1.65x, speed 1.15/s') alongside bounded 0-100% threat scores."),
        ("HUMAN-IN-THE-LOOP", "Decision-Support Telemetry", "Empowers coaches and safety supervisors with visual gauges and opt-in local logging. Autonomous enforcement is strictly prohibited.")
    ]

    for idx, (tag, stitle, sdesc) in enumerate(sol_cards):
        row = idx // 3
        col = idx % 3
        cx = Inches(0.8) + col * Inches(3.95)
        cy = Inches(2.1) + row * Inches(2.35)
        add_card(s3, cx, cy, Inches(3.75), Inches(2.15), tag, "")
        
        tb = s3.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.55), Inches(3.35), Inches(1.4))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = stitle
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(12)
        pt.font.bold = True
        pt.font.color.rgb = WHITE
        pt.space_after = Pt(4)
        
        pd = tf.add_paragraph()
        pd.text = sdesc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = GREY

    # --------------------------------------------------------------------------
    # SLIDE 4 — TECHNICAL ARCHITECTURE
    # --------------------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_background(s4)
    add_header(s4, "Engineering Pipeline", "Technical Architecture: From Raw Pixels to Explainable Alerts")

    # Left: Flowchart Blocks
    flow_steps = [
        ("1. Optical Input", "Live Webcam (720p/1080p) or Prerecorded Video Clip"),
        ("2. Frame Normalization", "OpenCV Ingestion & Downsampling to 960x540 Buffer"),
        ("3. Pose Estimation", "MediaPipe Pose Estimator (Active Provider: MediaPipe CPU; QNN Planned)"),
        ("4. Landmark Extraction", "9 Normalized Keypoints with Visibility Confidence Gating (>0.40)"),
        ("5. Temporal Sliding Window", "12-Frame Deque Tracking Velocity & Scale Growth"),
        ("6. Kinematic Pattern Engine", "Rule Evaluator: Fall > Punch > Kick > Approach > Guard > Neutral"),
        ("7. Threat Fusion & Telemetry", "Graduated Threat Scoring (0.0 - 1.0) & Rationale Generation"),
        ("8. Output & Storage", "AEGIS CORE Dynamic HUD Display & Local JSONL Incident Log")
    ]

    for idx, (s_title, s_desc) in enumerate(flow_steps):
        fy = Inches(1.95) + idx * Inches(0.62)
        box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), fy, Inches(5.8), Inches(0.52))
        box.fill.solid()
        box.fill.fore_color.rgb = PANEL_BG
        box.line.color.rgb = PANEL_BORDER
        box.line.width = Pt(1)

        tb = s4.shapes.add_textbox(Inches(0.9), fy + Inches(0.04), Inches(5.6), Inches(0.44))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r1 = p.add_run()
        r1.text = f"{s_title}: "
        r1.font.bold = True
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = CYAN
        r2 = p.add_run()
        r2.text = s_desc
        r2.font.size = Pt(9)
        r2.font.color.rgb = WHITE

    # Right: Embedded UI Preview Card
    add_card(s4, Inches(6.9), Inches(1.95), Inches(5.633), Inches(4.9), "AEGIS CORE TELEMETRY INTERFACE", "")
    if PREVIEW_IMG.exists():
        s4.shapes.add_picture(str(PREVIEW_IMG), Inches(7.1), Inches(2.55), Inches(5.233), Inches(2.94))
    
    lbl_box = s4.shapes.add_textbox(Inches(7.1), Inches(5.6), Inches(5.233), Inches(1.1))
    tf_lbl = lbl_box.text_frame
    tf_lbl.word_wrap = True
    p_l1 = tf_lbl.paragraphs[0]
    p_l1.text = "Current Verified Engine: MediaPipe CPU (Active)"
    p_l1.font.bold = True
    p_l1.font.size = Pt(10.5)
    p_l1.font.color.rgb = CYAN
    p_l2 = tf_lbl.add_paragraph()
    p_l2.text = "Heads-up display running at 960x540. QNN Hexagon NPU execution is planned via ONNX/QNN pose adapter."
    p_l2.font.size = Pt(9)
    p_l2.font.color.rgb = GREY

    # --------------------------------------------------------------------------
    # SLIDE 5 — INNOVATION, PRIVACY & SAFETY
    # --------------------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_background(s5)
    add_header(s5, "Responsible AI Architecture", "Innovation, Privacy, and Responsible Safety Safeguards")

    inno_cards = [
        ("ZERO FACIAL RECOGNITION", "Biometric Privacy by Design", "Aegis does not perform face recognition, face identification, facial embeddings or biometric identity tracking. One coarse nose keypoint is used only for pose geometry."),
        ("VOLATILE FRAME MEMORY", "No Cloud Video Uploads", "All frame computations execute in volatile RAM on the local workstation. Video streams are analyzed frame-by-frame and immediately overwritten."),
        ("OPT-IN LOCAL LOGGING", "Operator-Controlled Storage", "Incident logging is strictly opt-in via manual toggle (R or Space). Log records reside exclusively on the host PC in artifacts/incidents.jsonl."),
        ("INTENT VS. MOTION", "Kinematic Signals, Not Guilt", "Aegis measures biomechanical movement patterns; it does not and cannot infer criminal intent, malice, or moral culpability. Human context is mandatory."),
        ("DECISION-SUPPORT ONLY", "Prohibition of Autonomous Action", "Outputs are engineered strictly for human review. Under no circumstances should Aegis trigger automated physical enforcement or disciplinary interlocks."),
        ("SYNTHETIC EVALUATION", "Camera-Free Demonstration", "A deterministic synthetic demonstration mode (--demo and --export-preview) enables comprehensive validation without exposing real users or spaces.")
    ]

    for idx, (tag, ititle, idesc) in enumerate(inno_cards):
        row = idx // 3
        col = idx % 3
        cx = Inches(0.8) + col * Inches(3.95)
        cy = Inches(2.1) + row * Inches(2.35)
        add_card(s5, cx, cy, Inches(3.75), Inches(2.15), tag, "")
        
        tb = s5.shapes.add_textbox(cx + Inches(0.2), cy + Inches(0.55), Inches(3.35), Inches(1.4))
        tf = tb.text_frame
        tf.word_wrap = True
        pt = tf.paragraphs[0]
        pt.text = ititle
        pt.font.name = "Segoe UI"
        pt.font.size = Pt(12)
        pt.font.bold = True
        pt.font.color.rgb = WHITE
        pt.space_after = Pt(4)
        
        pd = tf.add_paragraph()
        pd.text = idesc
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(9.5)
        pd.font.color.rgb = GREY

    # --------------------------------------------------------------------------
    # SLIDE 6 — SNAPDRAGON OPTIMIZATION PATH
    # --------------------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_background(s6)
    add_header(s6, "Hardware Acceleration Strategy", "Qualcomm Snapdragon Optimization & NPU Deployment Path (Planned)")

    snap_steps = [
        ("STAGE 1: FUNCTIONAL BASELINE", "Current MediaPipe CPU Pipeline", "The validated MVP runs MediaPipe CPU to establish ground-truth algorithmic correctness, temporal stability, and deterministic test suites."),
        ("STAGE 2: MODEL SELECTION", "Qualcomm AI Hub Model Catalog", "Select an edge-optimized pose model (e.g., YOLOv8n-pose or RTMPose) from Qualcomm AI Hub with proven low-latency edge topology."),
        ("STAGE 3: COMPILATION & QUANTIZATION", "INT8 / FP16 Hexagon Compilation", "Utilize Qualcomm AI Hub cloud toolchains to compile and quantize the model for the Hexagon NPU on Snapdragon X Elite/Plus platforms."),
        ("STAGE 4: RUNTIME ADAPTATION (PLANNED)", "ONNX/QNN Pose-Estimator Adapter", "Implement adapter that preprocesses frames, calls session.run(), converts outputs to normalized PoseFrame landmarks, and replaces MediaPipePoseEstimator via QNN."),
        ("STAGE 5: RIGOROUS BENCHMARKING (PLANNED)", "Identical-Video CPU vs QNN Benchmark", "Execute scripts/benchmark.py on an identical 1080p clip to measure real throughput (FPS), p95 latency reduction, and CPU offload.")
    ]

    card_h = Inches(0.85)
    card_gap_y = Inches(0.16)
    for idx, (stage_tag, stitle, sdesc) in enumerate(snap_steps):
        sy = Inches(2.05) + idx * (card_h + card_gap_y)
        card = add_card(s6, Inches(0.8), sy, Inches(11.733), card_h, "", "")
        
        tb = s6.shapes.add_textbox(Inches(1.0), sy + Inches(0.08), Inches(11.3), Inches(0.7))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        r_num = p1.add_run()
        r_num.text = f"[{stage_tag}]  "
        r_num.font.bold = True
        r_num.font.size = Pt(10)
        r_num.font.color.rgb = CYAN
        r_tit = p1.add_run()
        r_tit.text = stitle
        r_tit.font.bold = True
        r_tit.font.size = Pt(11)
        r_tit.font.color.rgb = WHITE
        
        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = GREY

    note_box = s6.shapes.add_textbox(Inches(0.8), Inches(6.6), Inches(11.733), Inches(0.5))
    tf_n = note_box.text_frame
    p_n = tf_n.paragraphs[0]
    p_n.text = "* Architectural Transparency Note: Supplying AEGIS_MODEL_PATH does not connect model outputs to pose estimation. Runtime diagnostics only verify provider availability. MediaPipe CPU is the only currently verified active pose provider; QNN is planned."
    p_n.font.size = Pt(8.5)
    p_n.font.italic = True
    p_n.font.color.rgb = CYAN_DIM

    # --------------------------------------------------------------------------
    # SLIDE 7 — EVIDENCE, IMPACT & ROADMAP
    # --------------------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_background(s7)
    add_header(s7, "Verification & Future Vision", "Evidence, Real-World Impact, and Project Roadmap")

    # Column 1: Current Evidence
    add_card(s7, Inches(0.8), Inches(2.1), Inches(3.7), Inches(4.7), "VERIFIED CURRENT EVIDENCE", "")
    tb_ev = s7.shapes.add_textbox(Inches(1.0), Inches(2.7), Inches(3.3), Inches(3.9))
    tf_ev = tb_ev.text_frame
    tf_ev.word_wrap = True
    ev_points = [
        ("Functional Live Pipeline", "Real-time webcam ingestion, landmark tracking, threat scoring, and HUD."),
        ("17 / 17 Tests Passing", "Deterministic pytest suite verifying heuristics, safety bounds, and deduplication."),
        ("Synthetic Demo Suite", "Complete camera-free demonstration and visual preview exporter."),
        ("Truthful Telemetry", "HUD reports MediaPipe CPU. QNN NPU execution is planned via pose adapter."),
        ("Public GitHub Repository", "github.com/aniketkapgate7-crypto/Aegis-Fight-Pattern-Analyzer")
    ]
    for idx, (t, d) in enumerate(ev_points):
        p = tf_ev.paragraphs[0] if idx == 0 else tf_ev.add_paragraph()
        p.text = f"• {t}: {d}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = GREY
        p.space_after = Pt(8)

    # Column 2: Roadmap
    add_card(s7, Inches(4.8), Inches(2.1), Inches(3.7), Inches(4.7), "DEVELOPMENT ROADMAP", "")
    tb_rd = s7.shapes.add_textbox(Inches(5.0), Inches(2.7), Inches(3.3), Inches(3.9))
    tf_rd = tb_rd.text_frame
    tf_rd.word_wrap = True
    rd_points = [
        ("QNN Pose Model Integration (Planned)", "Implement ONNX/QNN pose adapter replacing MediaPipe on Hexagon NPU."),
        ("Snapdragon PC Benchmarking", "Evaluate throughput, latency, and power on Snapdragon X Elite HP laptops."),
        ("Multi-Person Interaction", "Extend kinematic engine to multi-subject spatial distance tracking."),
        ("Curated Validation Dataset", "Benchmark against controlled athletic sparring clips for recall/precision metrics.")
    ]
    for idx, (t, d) in enumerate(rd_points):
        p = tf_rd.paragraphs[0] if idx == 0 else tf_rd.add_paragraph()
        p.text = f"• {t}: {d}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = GREY
        p.space_after = Pt(10)

    # Column 3: Value Statement
    add_card(s7, Inches(8.8), Inches(2.1), Inches(3.733), Inches(4.7), "CLOSING VALUE STATEMENT", "")
    tb_vs = s7.shapes.add_textbox(Inches(9.0), Inches(2.7), Inches(3.333), Inches(3.9))
    tf_vs = tb_vs.text_frame
    tf_vs.word_wrap = True
    p_v1 = tf_vs.paragraphs[0]
    p_v1.text = "Edge-Native Motion Intelligence"
    p_v1.font.bold = True
    p_v1.font.size = Pt(14)
    p_v1.font.color.rgb = CYAN
    p_v1.space_after = Pt(12)
    
    p_v2 = tf_vs.add_paragraph()
    p_v2.text = "Aegis proves that advanced physical motion intelligence can run fully on-device without compromising personal privacy, streaming raw video to cloud servers, or harvesting facial biometrics."
    p_v2.font.size = Pt(10)
    p_v2.font.color.rgb = WHITE
    p_v2.space_after = Pt(14)

    p_v3 = tf_vs.add_paragraph()
    p_v3.text = "Designed for Snapdragon-powered HP PCs, Aegis showcases the transformative potential of edge NPU acceleration for human-in-the-loop safety and athletic intelligence."
    p_v3.font.size = Pt(10)
    p_v3.font.color.rgb = GREY

    prs.save(str(output_path))
    print(f"Pitch Deck PPTX generated at: {output_path.resolve()}")


def generate_pitch_deck_pdf(output_path: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.pdfgen import canvas

    # 16:9 widescreen format (13.333 x 7.5 inches = 960 x 540 pt)
    w_pt, h_pt = 13.333 * 72, 7.5 * 72  # 960 x 540 pt
    c = canvas.Canvas(str(output_path), pagesize=(w_pt, h_pt))

    BG_COLOR = colors.HexColor("#070F14")
    PANEL_BG = colors.HexColor("#0E1E28")
    PANEL_BORDER = colors.HexColor("#1C3A4D")
    CYAN = colors.HexColor("#00E5FF")
    CYAN_DIM = colors.HexColor("#008CBA")
    WHITE = colors.HexColor("#F5F8FA")
    GREY = colors.HexColor("#A5B4BC")

    def draw_bg():
        c.setFillColor(BG_COLOR)
        c.rect(0, 0, w_pt, h_pt, fill=1, stroke=0)

    def draw_header(category: str, title: str):
        c.setFont("Helvetica-Bold", 10)
        c.setFillColor(CYAN)
        c.drawString(58, h_pt - 42, f"/// {category.upper()}")
        c.setFont("Helvetica-Bold", 19)
        c.setFillColor(WHITE)
        c.drawString(58, h_pt - 68, title)
        c.setStrokeColor(PANEL_BORDER)
        c.setLineWidth(1)
        c.line(58, h_pt - 82, w_pt - 58, h_pt - 82)

    def draw_card(x, y, w, h, title=""):
        c.setFillColor(PANEL_BG)
        c.setStrokeColor(PANEL_BORDER)
        c.setLineWidth(1)
        c.roundRect(x, y, w, h, 6, fill=1, stroke=1)
        if title:
            c.setFont("Helvetica-Bold", 10.5)
            c.setFillColor(CYAN)
            c.drawString(x + 14, y + h - 22, title)

    # --------------------------------------------------------------------------
    # SLIDE 1
    # --------------------------------------------------------------------------
    draw_bg()
    c.setFillColor(CYAN)
    c.rect(80, h_pt - 360, 6, 220, fill=1, stroke=0)

    c.setFont("Helvetica-Bold", 36)
    c.setFillColor(CYAN)
    c.drawString(104, h_pt - 190, "AEGIS")

    c.setFont("Helvetica-Bold", 22)
    c.setFillColor(WHITE)
    c.drawString(104, h_pt - 230, "Privacy-Preserving On-Device Combat Motion Intelligence")

    c.setFont("Helvetica", 13)
    c.setFillColor(GREY)
    c.drawString(104, h_pt - 265, "Explainable real-time motion awareness designed for Snapdragon-powered HP PCs")

    c.setFont("Helvetica-Bold", 11)
    c.setFillColor(CYAN_DIM)
    c.drawString(104, h_pt - 320, "Participant: Aniket Kapgate  |  Snapdragon AI Lab Build & Present Challenge")
    c.showPage()

    # --------------------------------------------------------------------------
    # SLIDE 2
    # --------------------------------------------------------------------------
    draw_bg()
    draw_header("Challenge & Market Need", "The Problem: Latency, Fragility, and Surveillance Exposure")
    problems = [
        ("01 / REACTIVE REVIEW", "Incident Review Happens After the Fact", "In athletic coaching and facility safety, video is examined hours or days later. Supervisors lack real-time warning to intervene before injuries occur."),
        ("02 / CLOUD DEPENDENCE", "High Latency & Single Point of Failure", "Cloud video analytics streams video over public internet, introducing 200-800ms delays and failing completely during network interruptions."),
        ("03 / TEMPORAL BLINDNESS", "Single-Frame Models Miss Motion Context", "Static single-frame object detectors cannot distinguish a high-velocity punch from an open high-five or a stumble from a dynamic athletic slide."),
        ("04 / PRIVACY VULNERABILITY", "Biometric Harvesting & Facial Exposure", "Streaming surveillance footage to external cloud servers creates catastrophic risks of non-consensual biometric data scraping and surveillance abuses.")
    ]
    card_w = 198
    for i, (tag, t, d) in enumerate(problems):
        cx = 58 + i * 216
        draw_card(cx, 80, card_w, 360, tag)
        c.setFont("Helvetica-Bold", 11)
        c.setFillColor(WHITE)
        # title wrapping
        words = t.split()
        l1 = " ".join(words[:3])
        l2 = " ".join(words[3:])
        c.drawString(cx + 14, 395, l1)
        c.drawString(cx + 14, 380, l2)
        
        c.setFont("Helvetica", 8.8)
        c.setFillColor(GREY)
        # description lines
        d_words = d.split()
        cur_line = []
        cur_y = 350
        for w in d_words:
            cur_line.append(w)
            if len(" ".join(cur_line)) > 26:
                c.drawString(cx + 14, cur_y, " ".join(cur_line[:-1]))
                cur_line = [w]
                cur_y -= 14
        if cur_line:
            c.drawString(cx + 14, cur_y, " ".join(cur_line))
    c.showPage()

    # --------------------------------------------------------------------------
    # SLIDE 3
    # --------------------------------------------------------------------------
    draw_bg()
    draw_header("Product Vision", "The Solution: Edge-Native Biomechanical Intelligence")
    sol_cards = [
        ("LOCAL PROCESSING", "100% On-Device Execution", "Processes video frames entirely in local RAM on Snapdragon PCs. Video buffers are overwritten immediately."),
        ("ANONYMIZED SKELETONS", "Human Pose Landmarks", "Extracts 9 normalized skeletal landmarks. Strips identity: no face recognition or embeddings; nose keypoint used only for geometry."),
        ("TEMPORAL WINDOWS", "Multi-Frame Motion Context", "Maintains a 12-frame sliding window to calculate true physical velocities, body scale growth rates, and torso ratios."),
        ("ACTION VOCABULARY", "6 Biomechanical Patterns", "Deterministically detects Neutral, Guard, Punch-like extension, Kick-like extension, Rapid Approach, and Fall postures."),
        ("TRANSPARENT METRICS", "Explainable Telemetry", "Outputs human-verifiable evidence strings (e.g., 'right arm reach 1.65x, speed 1.15/s') alongside bounded threat scores."),
        ("HUMAN-IN-THE-LOOP", "Decision-Support Telemetry", "Empowers coaches and safety supervisors with visual gauges and opt-in local logging. Autonomous action is prohibited.")
    ]
    card_w3 = 270
    card_h3 = 165
    for i, (tag, t, d) in enumerate(sol_cards):
        r, col = i // 3, i % 3
        cx = 58 + col * 287
        cy = 265 - r * 185
        draw_card(cx, cy, card_w3, card_h3, tag)
        c.setFont("Helvetica-Bold", 10.5)
        c.setFillColor(WHITE)
        c.drawString(cx + 14, cy + card_h3 - 42, t)
        c.setFont("Helvetica", 8.5)
        c.setFillColor(GREY)
        d_words = d.split()
        cur_line = []
        cur_y = cy + card_h3 - 62
        for w in d_words:
            cur_line.append(w)
            if len(" ".join(cur_line)) > 38:
                c.drawString(cx + 14, cur_y, " ".join(cur_line[:-1]))
                cur_line = [w]
                cur_y -= 13
        if cur_line:
            c.drawString(cx + 14, cur_y, " ".join(cur_line))
    c.showPage()

    # --------------------------------------------------------------------------
    # SLIDE 4
    # --------------------------------------------------------------------------
    draw_bg()
    draw_header("Engineering Pipeline", "Technical Architecture: From Raw Pixels to Explainable Alerts")
    flow_steps = [
        ("1. Optical Input", "Webcam / Prerecorded Video Clip"),
        ("2. Normalization", "OpenCV Ingestion (960x540 Buffer)"),
        ("3. Pose Estimator", "MediaPipe Pose (Active: MediaPipe CPU; QNN Planned)"),
        ("4. Landmark Array", "9 Normalized Keypoints (Visibility > 0.40)"),
        ("5. Temporal Window", "12-Frame Deque Tracking Velocity"),
        ("6. Pattern Engine", "Kinematic Rule Evaluator (Priority Order)"),
        ("7. Threat Fusion", "Graduated Threat Index (0.0 - 1.0)"),
        ("8. Output & Logs", "AEGIS CORE HUD & Local JSONL Log")
    ]
    for i, (st, sd) in enumerate(flow_steps):
        fy = 395 - i * 44
        draw_card(58, fy, 410, 38)
        c.setFont("Helvetica-Bold", 9)
        c.setFillColor(CYAN)
        c.drawString(72, fy + 14, f"{st}: ")
        c.setFont("Helvetica", 8.5)
        c.setFillColor(WHITE)
        c.drawString(175, fy + 14, sd)

    draw_card(490, 85, 412, 350, "AEGIS CORE TELEMETRY INTERFACE")
    if PREVIEW_IMG.exists():
        c.drawImage(str(PREVIEW_IMG), 505, 150, width=382, height=215)
    c.setFont("Helvetica-Bold", 10)
    c.setFillColor(CYAN)
    c.drawString(505, 125, "Current Verified Engine: MediaPipe CPU (Active)")
    c.setFont("Helvetica", 8.5)
    c.setFillColor(GREY)
    c.drawString(505, 105, "Live HUD running at 960x540. QNN NPU execution is planned via pose adapter.")
    c.showPage()

    # --------------------------------------------------------------------------
    # SLIDE 5
    # --------------------------------------------------------------------------
    draw_bg()
    draw_header("Responsible AI Architecture", "Innovation, Privacy, and Responsible Safety Safeguards")
    inno_cards = [
        ("ZERO FACIAL RECOGNITION", "Biometric Privacy by Design", "Aegis does not perform face recognition, face identification, facial embeddings or biometric identity tracking. One coarse nose keypoint is used only for pose geometry."),
        ("VOLATILE FRAME MEMORY", "No Cloud Video Uploads", "All computations execute in volatile RAM on the local workstation. Video streams are never sent to external servers."),
        ("OPT-IN LOCAL LOGGING", "Operator-Controlled Storage", "Incident logging is strictly opt-in via manual toggle (R or Space). Log records reside exclusively on the host PC in artifacts/."),
        ("INTENT VS. MOTION", "Kinematic Signals, Not Guilt", "Aegis measures biomechanical movement patterns; it does not and cannot infer criminal intent or moral culpability."),
        ("DECISION-SUPPORT ONLY", "Prohibition of Autonomous Action", "Outputs are engineered strictly for human review. Under no circumstances should Aegis trigger automated punitive actions."),
        ("SYNTHETIC EVALUATION", "Camera-Free Demonstration", "A deterministic synthetic demonstration mode (--demo and --export-preview) enables comprehensive validation without real users.")
    ]
    for i, (tag, t, d) in enumerate(inno_cards):
        r, col = i // 3, i % 3
        cx = 58 + col * 287
        cy = 265 - r * 185
        draw_card(cx, cy, card_w3, card_h3, tag)
        c.setFont("Helvetica-Bold", 10.5)
        c.setFillColor(WHITE)
        c.drawString(cx + 14, cy + card_h3 - 42, t)
        c.setFont("Helvetica", 8.5)
        c.setFillColor(GREY)
        d_words = d.split()
        cur_line = []
        cur_y = cy + card_h3 - 62
        for w in d_words:
            cur_line.append(w)
            if len(" ".join(cur_line)) > 38:
                c.drawString(cx + 14, cur_y, " ".join(cur_line[:-1]))
                cur_line = [w]
                cur_y -= 13
        if cur_line:
            c.drawString(cx + 14, cur_y, " ".join(cur_line))
    c.showPage()

    # --------------------------------------------------------------------------
    # SLIDE 6
    # --------------------------------------------------------------------------
    draw_bg()
    draw_header("Hardware Acceleration Strategy", "Qualcomm Snapdragon Optimization & NPU Deployment Path (Planned)")
    snap_steps = [
        ("STAGE 1: FUNCTIONAL BASELINE", "Current MediaPipe CPU Pipeline", "The validated MVP runs MediaPipe CPU to establish ground-truth algorithmic correctness, temporal stability, and deterministic test suites."),
        ("STAGE 2: MODEL SELECTION", "Qualcomm AI Hub Model Catalog", "Select an edge-optimized pose model (e.g., YOLOv8n-pose or RTMPose) from Qualcomm AI Hub with proven low-latency edge topology."),
        ("STAGE 3: COMPILATION & QUANTIZATION", "INT8 / FP16 Hexagon Compilation", "Utilize Qualcomm AI Hub cloud toolchains to compile and quantize the model for the Hexagon NPU on Snapdragon X Elite/Plus platforms."),
        ("STAGE 4: ADAPTATION (PLANNED)", "ONNX/QNN Pose-Estimator Adapter", "Adapter preprocesses frames, calls session.run(), converts outputs to PoseFrame, and replaces MediaPipe."),
        ("STAGE 5: BENCHMARK (PLANNED)", "Identical-Video CPU vs QNN Benchmark", "Execute scripts/benchmark.py on an identical 1080p clip to measure real throughput (FPS) and p95 latency.")
    ]
    card_h6 = 58
    for i, (tag, t, d) in enumerate(snap_steps):
        sy = 370 - i * 70
        draw_card(58, sy, 844, card_h6)
        c.setFont("Helvetica-Bold", 9.5)
        c.setFillColor(CYAN)
        c.drawString(72, sy + 38, f"[{tag}]  ")
        c.setFillColor(WHITE)
        c.drawString(275, sy + 38, t)
        c.setFont("Helvetica", 8.5)
        c.setFillColor(GREY)
        c.drawString(72, sy + 18, d)

    c.setFont("Helvetica-Oblique", 8)
    c.setFillColor(CYAN_DIM)
    c.drawString(58, 48, "* Note: Supplying AEGIS_MODEL_PATH does not connect outputs to pose estimation. Diagnostics only verify provider availability. MediaPipe CPU is active; QNN is planned.")
    c.showPage()

    # --------------------------------------------------------------------------
    # SLIDE 7
    # --------------------------------------------------------------------------
    draw_bg()
    draw_header("Verification & Future Vision", "Evidence, Real-World Impact, and Project Roadmap")
    draw_card(58, 80, 268, 360, "VERIFIED CURRENT EVIDENCE")
    ev_lines = [
        "• Functional Live Pipeline: Real-time webcam ingestion, tracking, threat scoring.",
        "• 17 / 17 Tests Passing: Deterministic pytest suite verifying heuristics and safety bounds.",
        "• Synthetic Demo Suite: Complete camera-free demo & preview image exporter.",
        "• Truthful Telemetry: Truthfully reports MediaPipe CPU provider; QNN NPU execution is planned.",
        "• Public GitHub Repo: github.com/aniketkapgate7-crypto/Aegis-Fight-Pattern-Analyzer"
    ]
    ey = 385
    c.setFont("Helvetica", 8.5)
    c.setFillColor(GREY)
    for l in ev_lines:
        words = l.split()
        cline = []
        for w in words:
            cline.append(w)
            if len(" ".join(cline)) > 36:
                c.drawString(72, ey, " ".join(cline[:-1]))
                cline = [w]
                ey -= 13
        if cline:
            c.drawString(72, ey, " ".join(cline))
            ey -= 18

    draw_card(346, 80, 268, 360, "DEVELOPMENT ROADMAP")
    rd_lines = [
        "• QNN Pose Model (Planned): Implement ONNX/QNN pose adapter replacing MediaPipe on Hexagon NPU.",
        "• Snapdragon PC Benchmarks: Evaluate throughput, latency, and power on Snapdragon X Elite laptops.",
        "• Multi-Person Tracking: Extend kinematic engine to multi-subject spatial distance tracking.",
        "• Curated Validation Set: Benchmark against controlled athletic sparring clips for recall/precision."
    ]
    ry = 385
    c.setFont("Helvetica", 8.5)
    c.setFillColor(GREY)
    for l in rd_lines:
        words = l.split()
        cline = []
        for w in words:
            cline.append(w)
            if len(" ".join(cline)) > 36:
                c.drawString(360, ry, " ".join(cline[:-1]))
                cline = [w]
                ry -= 13
        if cline:
            c.drawString(360, ry, " ".join(cline))
            ry -= 18

    draw_card(634, 80, 268, 360, "CLOSING VALUE STATEMENT")
    c.setFont("Helvetica-Bold", 13)
    c.setFillColor(CYAN)
    c.drawString(648, 385, "Edge-Native Motion Intelligence")
    c.setFont("Helvetica", 9)
    c.setFillColor(WHITE)
    c.drawString(648, 350, "Aegis proves that advanced physical motion")
    c.drawString(648, 335, "intelligence can run fully on-device without")
    c.drawString(648, 320, "compromising personal privacy, streaming raw")
    c.drawString(648, 305, "video to cloud servers, or harvesting biometrics.")

    c.setFont("Helvetica", 9)
    c.setFillColor(GREY)
    c.drawString(648, 260, "Designed for Snapdragon-powered HP PCs,")
    c.drawString(648, 245, "Aegis showcases the transformative potential")
    c.drawString(648, 230, "of edge NPU acceleration for human-in-the-loop")
    c.drawString(648, 215, "safety and athletic intelligence.")
    c.showPage()

    c.save()
    print(f"Pitch Deck PDF generated at: {output_path.resolve()}")


def main() -> None:
    print("Generating submission materials...")
    
    # 1. Project Description
    docx_path = SUBMISSION_DIR / "Aegis_Project_Description.docx"
    pdf_desc_path = SUBMISSION_DIR / "Aegis_Project_Description.pdf"
    generate_project_description_docx(docx_path)
    generate_project_description_pdf(docx_path, pdf_desc_path)

    # 2. Pitch Deck
    pptx_path = SUBMISSION_DIR / "Aegis_Pitch_Deck.pptx"
    pdf_deck_path = SUBMISSION_DIR / "Aegis_Pitch_Deck.pdf"
    generate_pitch_deck_pptx(pptx_path)
    generate_pitch_deck_pdf(pdf_deck_path)

    print("All submission artifacts generated successfully!")


if __name__ == "__main__":
    main()
